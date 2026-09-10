import io
from datetime import date

import pytest
from openpyxl import Workbook, load_workbook

import interop_report as ir


def _summary_xlsx(rows, filename: str = "in.xlsx") -> tuple[str, bytes]:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["NO", "BANK_ID", "BANK_NAME", "ISSUER_TXN_COUNT",
               "ISSUER_TOTAL_AMOUNT", "ACQUIRER_TXN_COUNT", "ACQUIRER_TOTAL_AMOUNT"])
    for row in rows:
        ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    return filename, buf.getvalue()


# Banks are deliberately written with input-style names so alias mapping is
# exercised; all three resolve to QR canonical names.
_SAMPLE = [
    [1, "18", "Abay Bank", 92, 423763, 97, 765485],
    [2, "1", "Amhara Bank", 101, 2246260.76, 23, 1289245],
    [3, "26", "ZamZam", 10, 500.5, 30, 1000.25],
]


@pytest.mark.parametrize("series,raw,expected", [
    ("qr", "Abay Bank", "Abay"),
    ("qr", "Awash Bank", "Awash"),
    ("qr", "Coop", "CBO"),
    ("qr", "Buna Bank", "Bunna"),
    ("qr", "KAAFI", "KAAFI"),
    ("qr", "M-PESA", "M-PESA"),
    ("qr", "Nib Bank", "NIB"),
    ("p2p", "Abay Bank", "Abay"),
    ("p2p", "Buna Bank", "Buna Bank"),
    ("p2p", "KAAFI", "KAAFI MFI"),
    ("p2p", "MPESA", "MPESA"),
    ("p2p", "NibE Birr", "NibE Birr"),
    ("p2p", "Zemen Bank", "Zemen"),
    ("p2p", "Dedebit", "Dedebit MFI"),
])
def test_normalize_bank_name(series, raw, expected):
    assert ir.normalize_bank_name(raw, series) == expected


def test_parse_bank_summary_file():
    filename, data = _summary_xlsx(_SAMPLE)
    rows = ir.parse_bank_summary_file(data, filename)
    assert len(rows) == 3
    row = rows[0]
    assert row["BANK_NAME"] == "Abay Bank"
    assert row["ISSUER_TXN_COUNT"] == 92
    assert row["ISSUER_TOTAL_AMOUNT"] == 423763
    assert row["ACQUIRER_TXN_COUNT"] == 97
    assert row["ACQUIRER_TOTAL_AMOUNT"] == 765485


def test_parse_bank_summary_file_tolerant_headers():
    wb = Workbook()
    ws = wb.active
    ws.append(["NO", "BANK_ID", "BANK NAME", "Issuer TXN Count",
               "Issuer Total Amount", "Acquirer Txn Count", "Acquirer Total Amount"])
    ws.append([1, "18", "Abay Bank", 92, 423763, 97, 765485])
    buf = io.BytesIO()
    wb.save(buf)
    rows = ir.parse_bank_summary_file(buf.getvalue())
    assert rows[0]["BANK_NAME"] == "Abay Bank"
    assert rows[0]["ISSUER_TXN_COUNT"] == 92


def test_parse_rejects_file_without_summary_headers():
    wb = Workbook()
    ws = wb.active
    ws.append(["DESTINATION_BANK", "SOURCE_BANK", "AMOUNT"])
    ws.append(["A", "B", 100])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ValueError, match="could not locate"):
        ir.parse_bank_summary_file(buf.getvalue())


def test_merge_single_file_rows_and_total():
    filename, data = _summary_xlsx(_SAMPLE)
    records, per_file, warnings = ir.merge_bank_summary_files([(filename, data)], "qr")
    assert per_file == [{"filename": filename, "status": "ok", "rows": 3}]
    assert warnings == []
    assert len(records) == len(ir.QR_BANK_ORDER) + 1
    # canonical ordering: Abay is first, ZamZam is last
    assert records[0]["BANK_NAME"] == "Abay"
    assert records[-2]["BANK_NAME"] == "ZamZam"
    total = records[-1]
    assert total["BANK_NAME"] == "Total"
    assert total["ISSUER_TXN_COUNT"] == 92 + 101 + 10
    assert total["ISSUER_TOTAL_AMOUNT"] == round(423763 + 2246260.76 + 500.5, 2)
    assert total["ACQUIRER_TXN_COUNT"] == 97 + 23 + 30
    assert total["ACQUIRER_TOTAL_AMOUNT"] == round(765485 + 1289245 + 1000.25, 2)


def test_merge_across_multiple_files_sums_per_bank():
    f1, d1 = _summary_xlsx(_SAMPLE[:2], "a.xlsx")
    f2, d2 = _summary_xlsx(_SAMPLE[2:], "b.xlsx")
    records_a, _, _ = ir.merge_bank_summary_files([(f1, d1)], "qr")
    records_both, _, _ = ir.merge_bank_summary_files([(f1, d1), (f2, d2)], "qr")
    # records_a[0] is Abay with the first file's numbers only
    single_abay = next(r for r in records_a if r["BANK_NAME"] == "Abay")
    merged_abay = next(r for r in records_both if r["BANK_NAME"] == "Abay")
    assert merged_abay["ISSUER_TXN_COUNT"] == single_abay["ISSUER_TXN_COUNT"]
    assert merged_abay["ACQUIRER_TOTAL_AMOUNT"] == single_abay["ACQUIRER_TOTAL_AMOUNT"]
    merged_zamzam = next(r for r in records_both if r["BANK_NAME"] == "ZamZam")
    assert merged_zamzam["ISSUER_TXN_COUNT"] == 10
    assert merged_zamzam["ISSUER_TOTAL_AMOUNT"] == 500.5
    total = records_both[-1]
    assert total["BANK_NAME"] == "Total"
    assert total["ISSUER_TXN_COUNT"] == 92 + 101 + 10


def test_merge_with_missing_banks_leaves_zero_rows():
    # only two banks present -> the other canonical banks keep 0 values
    filename, data = _summary_xlsx(_SAMPLE[:2], "two.xlsx")
    records, _, _ = ir.merge_bank_summary_files([(filename, data)], "qr")
    present = [r["BANK_NAME"] for r in records if r["BANK_NAME"] not in ("Total", "TOTAL")]
    assert "Abay" in present and "Amhara" in present
    empty = next(r for r in records if r["BANK_NAME"] == "Ahadu Ebirr")
    assert empty["ISSUER_TXN_COUNT"] == 0
    assert empty["ISSUER_TOTAL_AMOUNT"] == 0.0


def test_merge_warns_on_unrecognised_bank():
    filename, data = _summary_xlsx([[9, "99", "Totally Unknown Bank", 1, 2.0, 3, 4.0]], "bad.xlsx")
    records, per_file, warnings = ir.merge_bank_summary_files([(filename, data)], "qr")
    assert per_file[0]["rows"] == 0
    assert any("unrecognised bank" in w for w in warnings)


def test_merge_corrupt_file_is_skipped():
    filename, data = _summary_xlsx(_SAMPLE, "ok.xlsx")
    records, per_file, warnings = ir.merge_bank_summary_files(
        [("broken.xlsx", b"not an xlsx"), (filename, data)], "qr")
    assert per_file[0]["status"] == "error"
    assert per_file[1]["status"] == "ok"
    assert len(warnings) == 1
    assert len(records) == len(ir.QR_BANK_ORDER) + 1


def test_p2p_series_keeps_ips_canonical_order():
    filename, data = _summary_xlsx(
        [[1, "21", "Zemen Bank", 10, 500, 20, 700],
         [2, "1", "Abay Bank", 3, 100, 4, 200]],
        "ips.xlsx",
    )
    records, _, _ = ir.merge_bank_summary_files([(filename, data)], "p2p")
    names = [r["BANK_NAME"] for r in records]
    assert names[:3] == ["Abay", "Abyssinia", "Addis"]
    assert "Zemen" in names
    assert records[-2]["BANK_NAME"] == "Zemen"
    assert records[-1]["BANK_NAME"] == "Total"
    assert len(records) == len(ir.IPS_BANK_ORDER) + 1


def test_success_report_filename():
    assert ir.success_report_filename("qr", date(2026, 9, 9)) == \
        "Successful QR Transaction for September 9,2026.xlsx"
    assert ir.success_report_filename("p2p", date(2026, 9, 9)) == \
        "Successful IPS Transaction for September 9,2026.xlsx"


def test_build_success_report_qr_structure():
    filename, data = _summary_xlsx([
        [1, "18", "Abay Bank", 92, 423763, 97, 765485],
        [7, "1", "Awash Bank", 0, 0, 23, 0],
        [36, "26", "ZamZam", 10, 0, 0, 1000],
    ], "qr.xlsx")
    records, _, _ = ir.merge_bank_summary_files([(filename, data)], "qr")
    out = ir.build_success_report_excel(records, "qr", date(2026, 9, 9))
    ws = load_workbook(io.BytesIO(out)).active
    assert ws.title == "Sheet1"
    # title block
    assert ws["C2"].value == "EthSwitch S.C."
    assert ws["C3"].value == "QR Report"
    assert "Successful QR Interoperable Transactions for September 09,2026" in ws["C5"].value
    # header block rows 6-7
    assert ws["C6"].value == "Bank"
    assert ws["D6"].value == "As a Destination"
    assert ws["F6"].value == "As a Source"
    assert ws["D7"].value == "No.Transactions"
    assert ws["E7"].value == "Values"
    # merges
    merged = {str(r) for r in ws.merged_cells.ranges}
    assert "C2:G2" in merged and "C6:C7" in merged and "D6:E6" in merged
    # data: Abay at row 8, Awash at row 14 (index 7), ZamZam at row 43
    assert ws["C8"].value == "Abay"
    assert ws["D8"].value == 92
    assert ws["G8"].value == 765485
    assert ws["C14"].value == "Awash"
    assert ws["C43"].value == "ZamZam"
    # zeros are rendered blank, non-zero kept
    assert ws.cell(row=14, column=4).value is None      # Awash issuer count 0
    assert ws.cell(row=14, column=5).value is None      # Awash issuer value 0
    assert ws.cell(row=14, column=6).value == 23
    assert ws.cell(row=14, column=7).value is None
    assert ws.cell(row=43, column=4).value == 10
    assert ws.cell(row=43, column=5).value is None
    assert ws.cell(row=43, column=6).value is None
    assert ws.cell(row=43, column=7).value == 1000
    # total row
    assert ws["C44"].value == "Total"
    assert ws["C44"].font.bold is False
    assert ws["D44"].value == "=SUM(D8:D43)"
    assert ws["D44"].font.bold is True
    assert ws["G44"].value == "=SUM(G8:G43)"
    # accounting number format applied to values
    assert ws["D8"].number_format.startswith("_ * #,##0")


def test_build_success_report_p2p_structure():
    filename, data = _summary_xlsx(
        [[1, "18", "Abay Bank", 9614, 46071133.45, 55810, 169001490.98],
         [51, "21", "Zemen Bank", 1280, 17612433.15, 1837, 22671018.96]],
        "ips.xlsx",
    )
    records, _, _ = ir.merge_bank_summary_files([(filename, data)], "p2p")
    out = ir.build_success_report_excel(records, "p2p", date(2026, 9, 9))
    ws = load_workbook(io.BytesIO(out)).active
    assert ws["B1"].value == "EthSwitch S.C."
    assert ws["B2"].value == "IPS Successful Report"
    assert "Successful IPS Interoperable Transactions Held on September 09,2026" in ws["B4"].value
    # headers: two spaces in the destination label
    assert ws["C5"].value == "Successful  Transactions As Destination"
    assert ws["E5"].value == "Successful Transactions As Source"
    assert ws["B5"].value == "BANK"
    assert ws["C6"].value == "No.Transactions"
    # data + total
    assert ws["B7"].value == "Abay"
    assert ws["C7"].value == 9614
    assert ws["F7"].value == 169001490.98
    assert ws["B58"].value == "TOTAL"
    assert ws["B58"].font.bold is True
    assert ws["C58"].value == "=SUM(C7:C57)"
    assert ws["F58"].value == "=SUM(F7:F57)"
    # value cells are right-aligned
    assert ws["C7"].alignment.horizontal == "right"


def test_build_merged_summary_excel():
    filename, data = _summary_xlsx(_SAMPLE, "qr.xlsx")
    records, _, _ = ir.merge_bank_summary_files([(filename, data)], "qr")
    out = ir.build_merged_summary_excel(records, "qr")
    ws = load_workbook(io.BytesIO(out)).active
    assert ws.title == "Report"
    assert ws["A1"].value == "Report name:"
    assert ws["B1"].value == "SUCCESSFUL QR TRANSACTION SUMMARY"
    header = [ws.cell(row=2, column=c).value for c in range(1, 7)]
    assert header[0] == "NO" and header[1] == "BANK_NAME"
    assert ws["A3"].value == 1 and ws["B3"].value == "Abay"
    total_r = 2 + 1 + len(ir.QR_BANK_ORDER)  # header + banks, TOTAL below
    assert ws.cell(row=total_r, column=2).value == "TOTAL"
    assert ws.cell(row=total_r, column=3).value == f"=SUM(C3:C{total_r - 1})"
    assert ws.cell(row=total_r, column=5).value == f"=SUM(E3:E{total_r - 1})"