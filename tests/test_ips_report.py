import io

import pytest
from openpyxl import Workbook, load_workbook

from ips_report import (
    parse_ips_trx_date,
    _label_for_key,
    parse_ips_report,
    collect_ips_dates,
    filter_ips_by_dates,
    build_ips_workbook,
)
from merger import IPS_CANONICAL_COLUMNS, IPS_MODE


def _trx(day, hh, mm, ss, frac="850000000"):
    return f"{day:02d}-SEP-26 {hh:02d}.{mm:02d}.{ss:02d}.{frac} PM"


def _no_header_export() -> bytes:
    """Mimic the real workbook: an 'Export Worksheet' caption sheet, empty
    Sheet1/Sheet2 and the data in 'Sheet10' with no header row."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Export Worksheet"
    ws.cell(row=1, column=1, value="DESTINATION_BANK")

    wb.create_sheet("Sheet1")
    wb.create_sheet("Sheet2")

    data = wb.create_sheet("Sheet10")
    data.append([None] * 8)  # blank first row (as in the real sample)
    data.append(["Commercial Bank of Ethiopia", "Awash Bank", _trx(16, 3, 27, 38),
                 "014251698580200", "1000409796544", 50.0, "260916182768094", "PROCESSED"])
    data.append(["Dashen Bank", "Awash Bank", _trx(16, 4, 5, 1),
                 "014251698580201", "1000409796545", 100.0, "260916182768095", "DECLINED"])
    data.append(["Dashen Bank", "CBE", _trx(15, 9, 30, 12),
                 "014251698580202", "1000409796546", 25.5, "260916182768096", "PROCESSED"])
    data.append([None] * 8)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _header_export() -> bytes:
    """A second sheet layout with a header row and reordered columns."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Source_bank", "Destination_Bank", "TRX_DATE", "amount",
               "DBTR_ACCT", "CDTR_ACCT", "tx_id", "STATUS"])
    ws.append(["Awash Bank", "CBE", _trx(14, 1, 2, 3), 10.0,
               "111", "222", "333", "PROCESSED"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_trx_date():
    assert parse_ips_trx_date("16-SEP-26 03.27.38.850000000 PM") == ("2026-09-16", "16-SEP-26")
    assert parse_ips_trx_date("06-SEP-26 10.11.12.000000000 AM") == ("2026-09-06", "06-SEP-26")
    assert parse_ips_trx_date("01-Jan-2026 00.00.00.0 AM") == ("2026-01-01", "01-JAN-26")
    assert parse_ips_trx_date("") is None
    assert parse_ips_trx_date("not a date") is None
    assert _label_for_key("2026-09-16") == "16-SEP-26"


def test_parse_no_header_scans_all_sheets():
    report = parse_ips_report(_no_header_export(), "sample.xlsx")
    assert len(report["records"]) == 3
    assert report["columns"] == list(IPS_CANONICAL_COLUMNS)

    first = report["records"][0]
    assert first["DESTINATION_BANK"] == "Commercial Bank of Ethiopia"
    assert first["SOURCE_BANK"] == "Awash Bank"
    assert first["STATUS"] == "PROCESSED"
    assert first["TX_ID"] == "260916182768094"

    dates = {d["key"]: d["count"] for d in report["dates"]}
    assert dates == {"2026-09-15": 1, "2026-09-16": 2}
    # every sheet is reported, even the empty ones
    assert "Sheet1" in report["sheets"] and report["sheets"]["Sheet1"] == 0
    assert report["sheets"]["Sheet10"] == 3


def test_parse_header_row_and_reordered_columns():
    report = parse_ips_report(_header_export(), "header.xlsx")
    assert len(report["records"]) == 1
    rec = report["records"][0]
    assert rec["SOURCE_BANK"] == "Awash Bank"
    assert rec["DESTINATION_BANK"] == "CBE"
    assert rec["AMOUNT"] == 10.0
    assert rec["STATUS"] == "PROCESSED"


def test_collect_and_filter_dates():
    r1 = parse_ips_report(_no_header_export(), "a.xlsx")
    r2 = parse_ips_report(_header_export(), "b.xlsx")
    dates = collect_ips_dates([r1, r2])
    keys = {d["key"] for d in dates}
    assert keys == {"2026-09-14", "2026-09-15", "2026-09-16"}

    all_records = r1["records"] + r2["records"]
    filtered = filter_ips_by_dates(all_records, ["2026-09-16"])
    assert len(filtered) == 2
    assert all("16-SEP-26" in r["TRX_DATE"] for r in filtered)

    filtered_all = filter_ips_by_dates(all_records, [d["key"] for d in dates])
    assert len(filtered_all) == 4

    assert filter_ips_by_dates(all_records, []) == []


def test_build_ips_workbook_layout():
    report = parse_ips_report(_no_header_export(), "sample.xlsx")
    data = build_ips_workbook(report["records"], "2026-09-15", "2026-09-16")
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    header = [c.value for c in ws[1]]
    assert header == list(IPS_CANONICAL_COLUMNS)
    assert ws.max_row == 1 + len(report["records"])
    # rows are sorted by TRX_DATE so the 15th comes before the 16th
    assert "15-SEP-26" in ws.cell(row=2, column=3).value
    assert IPS_MODE.key == "ips"


# ── Flask endpoints ─────────────────────────────────────────────────────────


@pytest.fixture
def client():
    import app as app_module
    app_module.app.config["TESTING"] = True
    return app_module.app.test_client()


def test_flask_ips_analyze_and_merge(client):
    data = _no_header_export()
    resp = client.post(
        "/ips-analyze",
        data={"files": (io.BytesIO(data), "sept.xlsx")},
        content_type="multipart/form-data",
    )
    assert resp.status_code == 200
    payload = resp.get_json()
    assert payload["total_rows"] == 3
    dates = [d["key"] for d in payload["dates"]]
    assert dates == ["2026-09-15", "2026-09-16"]
    token = payload["token"]

    merged = client.post(
        "/ips-merge",
        data={"token": token, "dates": "2026-09-16"},
        content_type="multipart/form-data",
    )
    assert merged.status_code == 200
    m = merged.get_json()
    assert m["total_rows"] == 2
    assert m["mode"] == "ips"
    assert m["filename"].startswith("IPS_Transactions_16-SEP-26")
    assert m["columns"] == list(IPS_CANONICAL_COLUMNS)

    dl = client.get(f"/download/{m['token']}")
    assert dl.status_code == 200
    wb = load_workbook(io.BytesIO(dl.data))
    assert [c.value for c in wb.active[1]] == list(IPS_CANONICAL_COLUMNS)


def test_flask_ips_requires_dates(client):
    data = _no_header_export()
    resp = client.post(
        "/ips-analyze",
        data={"files": (io.BytesIO(data), "sept.xlsx")},
        content_type="multipart/form-data",
    )
    token = resp.get_json()["token"]
    bad = client.post(
        "/ips-merge",
        data={"token": token, "dates": "[]"},
        content_type="multipart/form-data",
    )
    assert bad.status_code == 400


def test_flask_ips_analyze_rejects_empty(client):
    data = _no_header_export()
    # an empty workbook has no IPS rows
    wb = Workbook()
    wb.active.title = "Sheet1"
    buf = io.BytesIO()
    wb.save(buf)
    resp = client.post(
        "/ips-analyze",
        data={"files": (io.BytesIO(buf.getvalue()), "empty.xlsx")},
        content_type="multipart/form-data",
    )
    assert resp.status_code == 400
    assert "No IPS transaction rows" in resp.get_json()["error"]
